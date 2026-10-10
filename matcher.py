import html
import re
from io import BytesIO
from urllib.parse import urlparse

import requests
from docx import Document
from pypdf import PdfReader


# ---------------------------------------------------------
# API CONFIGURATION
# ---------------------------------------------------------

HIMALAYAS_API = "https://himalayas.app/jobs/api/search"
ARBEITNOW_API = "https://www.arbeitnow.com/api/job-board-api"

HEADERS = {
    "User-Agent": "AI-Career-Agent/1.0"
}

TIMEOUT = 20


# ---------------------------------------------------------
# SKILL LIBRARY
# ---------------------------------------------------------

SKILL_LIBRARY = [
    "Excel",
    "Advanced Excel",
    "SQL",
    "Python",
    "Power BI",
    "Tableau",
    "Data Visualization",
    "Data Cleaning",
    "Statistics",
    "Machine Learning",
    "Pandas",
    "NumPy",
    "Scikit-learn",
    "Jupyter",
    "VLOOKUP",
    "XLOOKUP",
    "Pivot Tables",
    "Power Query",
    "Business Reporting",
    "Financial Analysis",
    "Dashboard",
    "Google Sheets",
    "R",
    "Looker",
    "ETL",
    "Data Analysis",
    "Tally",
    "Accounting",
    "Sales Analysis",
    "Data Modeling",
]


# ---------------------------------------------------------
# GENERAL HELPERS
# ---------------------------------------------------------

def clean_text(value):
    """Convert HTML and other values into readable text."""

    value = re.sub(
        r"<[^>]+>",
        " ",
        str(value or ""),
    )

    value = html.unescape(value)

    return re.sub(
        r"\s+",
        " ",
        value,
    ).strip()


def safe_url(value):
    """Return a URL only if it uses HTTP or HTTPS."""

    value = str(value or "").strip()

    if not value:
        return ""

    parsed = urlparse(value)

    if parsed.scheme in ("http", "https") and parsed.netloc:
        return value

    return ""


# =========================================================
# CV READING
# =========================================================

def extract_resume_text(uploaded_file):
    """
    Extract readable text from an uploaded PDF or DOCX file.
    """

    file_bytes = uploaded_file.getvalue()
    filename = uploaded_file.name.lower()

    # ---------------- PDF ----------------

    if filename.endswith(".pdf"):

        reader = PdfReader(
            BytesIO(file_bytes)
        )

        if reader.is_encrypted:
            raise ValueError(
                "This PDF is password-protected. "
                "Please upload an unlocked PDF."
            )

        text_parts = []

        for page in reader.pages:
            page_text = page.extract_text() or ""
            text_parts.append(page_text)

        text = "\n".join(text_parts)

    # ---------------- DOCX ----------------

    elif filename.endswith(".docx"):

        document = Document(
            BytesIO(file_bytes)
        )

        parts = []

        # Normal paragraphs
        for paragraph in document.paragraphs:
            if paragraph.text.strip():
                parts.append(paragraph.text)

        # Tables
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        parts.append(cell.text)

        text = "\n".join(parts)

    else:
        raise ValueError(
            "Please upload a PDF or DOCX file."
        )

    if not text.strip():
        raise ValueError(
            "No readable text was found in this file. "
            "If this is a scanned PDF, OCR support will be "
            "needed."
        )

    return text.strip()


# =========================================================
# CV SKILL EXTRACTION
# =========================================================

def extract_resume_skills(resume_text):
    """
    Identify recognised skills mentioned in the CV.
    """

    found = []

    text = resume_text.lower()

    for skill in SKILL_LIBRARY:

        pattern = (
            r"(?<!\w)"
            + re.escape(skill.lower())
            + r"(?!\w)"
        )

        if re.search(pattern, text):
            found.append(skill)

    # Remove duplicates while preserving order
    return list(dict.fromkeys(found))


# =========================================================
# EXPERIENCE ESTIMATION
# =========================================================

def estimate_experience(resume_text):
    """
    Estimate experience only when the CV explicitly
    states a number of years.
    """

    text = resume_text.lower()

    patterns = [
        r"(\d+(?:\.\d+)?)\+?\s+years?\s+(?:of\s+)?"
        r"(?:professional\s+)?experience",

        r"experience\s*:\s*"
        r"(\d+(?:\.\d+)?)\+?\s+years?",
    ]

    values = []

    for pattern in patterns:

        matches = re.findall(
            pattern,
            text,
        )

        values.extend(matches)

    if values:

        years = max(
            float(value)
            for value in values
        )

        return {
            "years": years,
            "label": (
                f"{years:g} years explicitly stated in CV"
            ),
            "confidence": (
                "Estimate from CV wording; "
                "verify manually."
            ),
        }

    return {
        "years": None,
        "label": (
            "Experience duration not identified"
        ),
        "confidence": (
            "Needs manual review."
        ),
    }


# =========================================================
# BUILD CV PROFILE
# =========================================================

def build_resume_profile(resume_text):
    """
    Build a structured profile from the CV.
    """

    return {
        "skills": extract_resume_skills(
            resume_text
        ),

        "experience": estimate_experience(
            resume_text
        ),

        "resume_text": resume_text,
    }


# =========================================================
# VISA / SPONSORSHIP CLASSIFICATION
# =========================================================

def classify_visa(text):
    """
    Classify sponsorship language found in the listing.

    IMPORTANT:
    This does NOT independently verify sponsorship.
    """

    text = clean_text(text).lower()

    negative_patterns = [

        r"\bno visa sponsorship\b",

        r"\bwill not sponsor\b",

        r"\bunable to sponsor\b",

        r"\bdoes not sponsor\b",

        r"\bwithout sponsorship\b",

        r"\bmust already have (?:the )?right to work\b",

        r"\bexisting work authorization required\b",

        r"\bnot able to provide sponsorship\b",
    ]

    positive_patterns = [

        r"\bvisa sponsorship (?:is )?available\b",

        r"\bvisa sponsorship provided\b",

        r"\bwill sponsor (?:a )?visa\b",

        r"\bsponsor(?:s|ing)? "
        r"(?:your |a |the )?visa\b",

        r"\bwork visa sponsorship\b",

        r"\bskilled worker visa sponsorship\b",

        r"\bvisa support provided\b",

        r"\brelocation and visa support\b",
    ]

    # Check negative language first
    if any(
        re.search(pattern, text)
        for pattern in negative_patterns
    ):
        return (
            "Not available / work authorization required"
        )

    # Check positive language
    if any(
        re.search(pattern, text)
        for pattern in positive_patterns
    ):
        return (
            "Sponsorship explicitly mentioned"
        )

    return (
        "Unclear - confirm with employer"
    )


def visa_evidence(text, visa_label):
    """
    Return the text evidence used for sponsorship
    classification.
    """

    text = clean_text(text)
    lower = text.lower()

    patterns = {

        "Sponsorship explicitly mentioned": [

            r"\bvisa sponsorship (?:is )?available\b",

            r"\bvisa sponsorship provided\b",

            r"\bwill sponsor (?:a )?visa\b",

            r"\bsponsor(?:s|ing)? "
            r"(?:your |a |the )?visa\b",

            r"\bwork visa sponsorship\b",

            r"\bskilled worker visa sponsorship\b",

            r"\bvisa support provided\b",

            r"\brelocation and visa support\b",
        ],

        "Not available / work authorization required": [

            r"\bno visa sponsorship\b",

            r"\bwill not sponsor\b",

            r"\bunable to sponsor\b",

            r"\bdoes not sponsor\b",

            r"\bwithout sponsorship\b",

            r"\bmust already have "
            r"(?:the )?right to work\b",

            r"\bexisting work authorization required\b",

            r"\bnot able to provide sponsorship\b",
        ],
    }

    for pattern in patterns.get(
        visa_label,
        [],
    ):

        match = re.search(
            pattern,
            lower,
        )

        if match:

            start = max(
                0,
                match.start() - 100,
            )

            end = min(
                len(text),
                match.end() + 140,
            )

            return {
                "evidence": text[start:end],

                "verification": (
                    "Keyword found in listing. "
                    "Verify sponsorship directly "
                    "with the employer."
                ),
            }

    return {
        "evidence": (
            "No clear sponsorship statement "
            "detected in the listing."
        ),

        "verification": (
            "Unverified — ask the employer directly."
        ),
    }


# =========================================================
# CONTACT DETAILS
# =========================================================

def extract_contacts(text):
    """
    Extract public email addresses and phone numbers
    only when they appear in the supplied listing text.
    """

    text = clean_text(text)

    emails = sorted(
        set(
            re.findall(
                r"\b[A-Z0-9._%+-]+"
                r"@[A-Z0-9.-]+\.[A-Z]{2,}\b",
                text,
                flags=re.I,
            )
        )
    )

    phones = sorted(
        set(
            re.findall(
                r"(?<!\w)"
                r"(?:\+\d{1,3}[\s().-]?)?"
                r"(?:\(\d{2,5}\)|\d{2,5})"
                r"[\s.-]\d{3,5}"
                r"[\s.-]\d{3,5}"
                r"(?!\w)",
                text,
            )
        )
    )

    return {
        "emails": emails[:5],
        "phones": phones[:5],
    }


# =========================================================
# SALARY
# =========================================================

def format_salary(job):

    low = job.get("minSalary")
    high = job.get("maxSalary")

    currency = (
        job.get("currency")
        or ""
    )

    period = (
        job.get("salaryPeriod")
        or "annual"
    )

    def display_amount(value):

        if value is None:
            return None

        try:
            return f"{float(value):,.0f}"

        except (
            ValueError,
            TypeError,
        ):
            return str(value)

    if low is None and high is None:
        return "Not listed"

    if (
        low is not None
        and high is not None
    ):

        amount = (
            f"{display_amount(low)}"
            f" - "
            f"{display_amount(high)}"
        )

    else:

        amount = display_amount(
            low
            if low is not None
            else high
        )

    return (
        f"{currency} {amount} ({period})"
    ).strip()


# =========================================================
# HIMALAYAS NORMALIZATION
# =========================================================

def normalize_himalayas(job):

    description = clean_text(
        job.get("description")
        or job.get("excerpt")
        or ""
    )

    # Location restrictions
    restrictions = (
        job.get("locationRestrictions")
        or []
    )

    if not isinstance(
        restrictions,
        list,
    ):
        restrictions = [
            restrictions
        ]

    restrictions = [

        (
            item.get("name", "")
            if isinstance(item, dict)
            else str(item)
        )

        for item in restrictions
    ]

    restrictions = [
        item
        for item in restrictions
        if item.strip()
    ]

    # Seniority
    seniority = (
        job.get("seniority")
        or []
    )

    if not isinstance(
        seniority,
        list,
    ):
        seniority = [
            seniority
        ]

    seniority_text = (
        ", ".join(
            map(str, seniority)
        )
        or "Not specified"
    )

    title = job.get(
        "title",
        "Untitled",
    )

    visa_text = (
        f"{title} {description}"
    )

    visa = classify_visa(
        visa_text
    )

    return {

        "title": title,

        "company": job.get(
            "companyName",
            "Not listed",
        ),

        "location": (
            ", ".join(restrictions)
            if restrictions
            else "Worldwide / not restricted"
        ),

        "country_restrictions": (
            restrictions
        ),

        "worldwide": not bool(
            restrictions
        ),

        "remote": True,

        "job_type": job.get(
            "employmentType",
            "Not listed",
        ),

        "experience": seniority_text,

        "salary": format_salary(
            job
        ),

        "description": description,

        "visa": visa,

        "visa_details": visa_evidence(
            visa_text,
            visa,
        ),

        "contact": extract_contacts(
            description
        ),

        "url": safe_url(
            job.get("applicationLink")
            or job.get("guid")
        ),

        "company_url": safe_url(
            job.get("companyWebsite")
            or job.get("companyUrl")
        ),

        "source": "Himalayas",
    }


# =========================================================
# ARBEITNOW NORMALIZATION
# =========================================================

def normalize_arbeitnow(job):

    description = clean_text(
        job.get(
            "description",
            "",
        )
    )

    remote_value = job.get(
        "remote",
        False,
    )

    is_remote = (
        remote_value is True
        or str(remote_value).lower()
        in (
            "true",
            "yes",
            "remote",
        )
    )

    location = job.get(
        "location",
        "Not listed",
    )

    title = job.get(
        "title",
        "Untitled",
    )

    visa_text = (
        f"{title} {description}"
    )

    visa = classify_visa(
        visa_text
    )

    job_types = (
        job.get("job_types")
        or []
    )

    if not isinstance(
        job_types,
        list,
    ):
        job_types = [
            job_types
        ]

    return {

        "title": title,

        "company": job.get(
            "company_name",
            "Not listed",
        ),

        "location": location,

        "country_restrictions": [],

        "worldwide": False,

        "remote": is_remote,

        "job_type": (
            ", ".join(
                map(
                    str,
                    job_types,
                )
            )
            or "Not listed"
        ),

        "experience": (
            "Not specified"
        ),

        "salary": "Not listed",

        "description": description,

        "visa": visa,

        "visa_details": visa_evidence(
            visa_text,
            visa,
        ),

        "contact": extract_contacts(
            description
        ),

        "url": safe_url(
            job.get("url")
        ),

        "company_url": safe_url(
            job.get("company_website")
            or job.get("company_url")
        ),

        "source": "Arbeitnow",
    }


# =========================================================
# JOB SEARCH
# =========================================================

def fetch_jobs(
    position,
    country,
    include_remote=True,
):

    jobs = []
    errors = []

    # -----------------------------------------------------
    # HIMALAYAS
    # -----------------------------------------------------

    try:

        params = {
            "q": position,
            "page": 1,
        }

        if country:
            params["country"] = country

        response = requests.get(
            HIMALAYAS_API,
            params=params,
            headers=HEADERS,
            timeout=TIMEOUT,
        )

        response.raise_for_status()

        data = response.json()

        for item in data.get(
            "jobs",
            [],
        ):

            jobs.append(
                normalize_himalayas(
                    item
                )
            )

    except (
        requests.RequestException,
        ValueError,
        TypeError,
    ) as exc:

        errors.append(
            f"Himalayas: {exc}"
        )

    # -----------------------------------------------------
    # ARBEITNOW
    # -----------------------------------------------------

    try:

        for page in range(
            1,
            4,
        ):

            response = requests.get(
                ARBEITNOW_API,
                params={
                    "page": page
                },
                headers=HEADERS,
                timeout=TIMEOUT,
            )

            response.raise_for_status()

            data = response.json().get(
                "data",
                [],
            )

            if not data:
                break

            for item in data:

                jobs.append(
                    normalize_arbeitnow(
                        item
                    )
                )

    except (
        requests.RequestException,
        ValueError,
        TypeError,
    ) as exc:

        errors.append(
            f"Arbeitnow: {exc}"
        )

    # -----------------------------------------------------
    # DEDUPLICATE
    # -----------------------------------------------------

    unique = {}

    for job in jobs:

        key = (
            job.get("url")
            or (
                job.get("company"),
                job.get("title"),
                job.get("location"),
            )
        )

        unique[key] = job

    # -----------------------------------------------------
    # ROLE + COUNTRY FILTER
    # -----------------------------------------------------

    country_lower = (
        str(country or "")
        .strip()
        .lower()
    )

    role_words = set(
        re.findall(
            r"\w+",
            str(position or "").lower(),
        )
    )

    filtered = []

    for job in unique.values():

        title_words = set(
            re.findall(
                r"\w+",
                job.get(
                    "title",
                    "",
                ).lower(),
            )
        )

        location = (
            job.get(
                "location",
                "",
            )
            .lower()
        )

        restrictions = [
            str(item).lower()
            for item in job.get(
                "country_restrictions",
                [],
            )
        ]

        # -------------------------------------------------
        # ROLE MATCH
        # -------------------------------------------------

        if role_words:

            if not (
                role_words
                & title_words
            ):
                continue

        # -------------------------------------------------
        # COUNTRY MATCH
        # -------------------------------------------------

        if country_lower:

            country_matches = (

                country_lower
                in location

                or any(
                    country_lower == item
                    or country_lower in item
                    for item in restrictions
                )
            )

            worldwide_remote = (
                include_remote
                and job.get(
                    "worldwide",
                    False,
                )
                and job.get(
                    "remote",
                    False,
                )
            )

            if not country_matches and not worldwide_remote:
                continue

        # -------------------------------------------------
        # REMOTE FILTER
        # -------------------------------------------------

        if (
            not include_remote
            and job.get(
                "remote",
                False,
            )
        ):
            continue

        filtered.append(
            job
        )

    return (
        filtered,
        errors,
    )


# =========================================================
# CV → JOB MATCHING
# =========================================================

def match_job_to_profile(
    job,
    resume_profile,
):
    """
    Compare CV skills with the job listing.

    This is intentionally a transparent keyword-based
    score, not an AI hiring prediction.
    """

    resume_profile = (
        resume_profile
        or {}
    )

    cv_skills = (
        resume_profile.get(
            "skills",
            [],
        )
    )

    job_text = (
        job.get(
            "title",
            "",
        )
        + " "
        + job.get(
            "description",
            "",
        )
    ).lower()

    # -----------------------------------------------------
    # MATCHED SKILLS
    # -----------------------------------------------------

    matched = []

    for skill in cv_skills:

        if skill.lower() in job_text:
            matched.append(
                skill
            )

    # -----------------------------------------------------
    # SKILLS NOT FOUND
    # -----------------------------------------------------

    missing = [

        skill

        for skill in cv_skills

        if skill.lower()
        not in job_text
    ]

    # -----------------------------------------------------
    # SCORE
    # -----------------------------------------------------

    if cv_skills:

        score = round(
            100
            * len(matched)
            / len(cv_skills)
        )

    else:
        score = 0

    # -----------------------------------------------------
    # SENIORITY CHECK
    # -----------------------------------------------------

    experience = (
        resume_profile.get(
            "experience",
            {},
        )
    )

    cv_years = experience.get(
        "years"
    )

    seniority_warning = None

    senior_patterns = [

        r"\bsenior\b",

        r"\blead\b",

        r"\bprincipal\b",

        r"\bstaff\b",

        r"\bmanager\b",

        r"\b\d+\+?\s+years?\s+"
        r"(?:of\s+)?experience\b",
    ]

    has_senior_language = any(
        re.search(
            pattern,
            job_text,
        )
        for pattern in senior_patterns
    )

    if (
        cv_years is not None
        and cv_years < 3
        and has_senior_language
    ):

        seniority_warning = (
            "Possible seniority mismatch: "
            "the listing mentions a senior title "
            "or experience requirement. "
            "Review it manually."
        )

    elif (
        cv_years is None
        and any(
            term in job_text
            for term in [
                "senior",
                "lead",
                "principal",
                "staff",
            ]
        )
    ):

        seniority_warning = (
            "Potentially senior role. "
            "Your CV experience duration "
            "could not be determined automatically."
        )

    return {

        "match_score": score,

        "matched_skills": matched,

        "missing_skills": missing,

        "seniority_warning": (
            seniority_warning
        ),

        "score_note": (
            "Keyword overlap only. "
            "Review the complete job description "
            "before applying."
        ),
    }


# =========================================================
# JOB FILTERING
# =========================================================

def filter_jobs(
    jobs,
    job_type,
    visa_filter,
    experience_filter,
    resume_profile=None,
):

    results = []

    for original_job in jobs:

        job = dict(
            original_job
        )

        job_text = (
            job.get(
                "title",
                "",
            )
            + " "
            + job.get(
                "description",
                "",
            )
        ).lower()

        # -------------------------------------------------
        # WORK ARRANGEMENT
        # -------------------------------------------------

        if (
            job_type == "Remote"
            and not job.get(
                "remote",
                False,
            )
        ):
            continue

        if (
            job_type == "On-site"
            and job.get(
                "remote",
                False,
            )
        ):
            continue

        if (
            job_type == "Hybrid"
            and "hybrid" not in job_text
        ):
            continue

        # -------------------------------------------------
        # VISA FILTER
        # -------------------------------------------------

        if (
            visa_filter
            == "Sponsorship explicitly mentioned"
        ):

            if job.get("visa") != (
                "Sponsorship explicitly mentioned"
            ):
                continue

        elif (
            visa_filter
            == "Sponsorship unclear"
        ):

            if job.get("visa") != (
                "Unclear - confirm with employer"
            ):
                continue

        elif (
            visa_filter
            == "No sponsorship / existing work permit"
        ):

            if job.get("visa") != (
                "Not available / work authorization required"
            ):
                continue

        # -------------------------------------------------
        # EXPERIENCE FILTER
        # -------------------------------------------------

        if (
            experience_filter
            == "Entry-level / 0-2 years"
        ):

            seniority = (
                job.get(
                    "experience",
                    "",
                )
                .lower()
                .strip()
            )

            entry_level_terms = [

                "entry",

                "junior",

                "intern",

                "graduate",

                "0-2",

                "0–2",

                "0 to 2",
            ]

            is_entry_level = (
                any(
                    term in seniority
                    for term
                    in entry_level_terms
                )

                or bool(
                    re.search(
                        r"\b(?:0|1|2)\s*"
                        r"(?:-|to)\s*\d+\s*"
                        r"years?\b",
                        job_text,
                    )
                )
            )

            seniority_unknown = (
                not seniority
                or seniority
                in (
                    "not specified",
                    "unknown",
                )
            )

            # Keep unknown listings.
            # Do not silently throw them away.
            if (
                not is_entry_level
                and not seniority_unknown
            ):
                continue

        # -------------------------------------------------
        # PERSONALIZED CV MATCH
        # -------------------------------------------------

        if resume_profile:

            job["profile_match"] = (
                match_job_to_profile(
                    job,
                    resume_profile,
                )
            )

        results.append(
            job
        )

    return results